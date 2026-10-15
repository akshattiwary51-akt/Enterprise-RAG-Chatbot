"""SYNTHETIC DEMO DATA -- development/testing only. All people, figures and 'secrets' are invented.

The CANARY_* strings are distinctive tokens: if one ever appears in a response for a user who is not
entitled to the document that contains it, access control is broken. Tests rely on this.
"""
from __future__ import annotations

import io
from pathlib import Path

from .core.security import hash_password
from .extensions import db
from .models import Document, DocumentRoleGrant, Role, User
from .services.ingestion import pipeline
from .services.ingestion.validation import validate_upload
from .services.permissions.policy import seed_roles_and_permissions

CANARY_HR = "CANARY-HR-7Q3X"
CANARY_FIN = "CANARY-FIN-9K2M"
CANARY_ENG = "CANARY-ENG-4D8P"
CANARY_PUBLIC = "CANARY-ALL-1A5B"

HR_POLICY_MD = f"""# Leave Policy (SYNTHETIC)

## Annual Leave
Full-time employees receive 25 days of paid annual leave per year. Leave must be approved by a manager.

## HR Confidential
The internal grievance escalation code is {CANARY_HR}. Do not share it outside the HR team.
"""

ENG_ARCH_MD = f"""# Payments Platform Architecture (SYNTHETIC)

## Overview
The payments platform uses an event-driven architecture with a message queue between the API gateway and the settlement service.

## Secrets Handling
The staging signing key identifier is {CANARY_ENG}. Rotate it every 90 days.
"""

HANDBOOK_TXT = f"""Company Handbook (SYNTHETIC)

Working hours are 9am to 5pm Monday to Friday. The office is closed on public holidays.
The all-hands meeting happens on the first Monday of every month. Reference: {CANARY_PUBLIC}.
"""

SALARY_ROWS = [
    ("Employee", "Department", "Role", "Salary"),
    ("Alice Example", "Engineering", "Senior Engineer", 120000),
    ("Bob Example", "Engineering", "Junior Developer", 80000),
    ("Carol Example", "HR", "HR Manager", 95000),
    ("Dan Example", "Finance", "Analyst", 70000),
    (f"Eve {CANARY_HR}", "HR", "Recruiter", 65000),
]

BUDGET_ROWS = [
    ("Quarter", "Cost Center", "Budget"),
    ("Q1", "Marketing", 50000),
    ("Q1", "Engineering", 150000),
    ("Q2", "Marketing", 55000),
    ("Q2", "Engineering", 160000),
    ("Q2", f"Reserve {CANARY_FIN}", 12345),
]


def make_xlsx(rows, sheet="Sheet1") -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_csv(rows) -> bytes:
    return ("\n".join(",".join(str(c) for c in r) for r in rows) + "\n").encode()


def make_docx(paragraphs: list[tuple[str, str]]) -> bytes:
    """paragraphs: [(style, text)] with style 'Heading 1', 'Heading 2' or 'Normal'."""
    from docx import Document as Docx

    d = Docx()
    for style, text in paragraphs:
        if style.startswith("Heading"):
            d.add_heading(text, level=int(style.split()[-1]))
        else:
            d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf(pages: list[str]) -> bytes:
    """Minimal valid multi-page text PDF (Helvetica) -- no external dependency."""
    objs: list[bytes] = []

    def add(b: bytes) -> int:
        objs.append(b)
        return len(objs)

    add(b"<< /Type /Catalog /Pages 2 0 R >>")
    add(b"")  # pages placeholder (object 2)
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    kids = []
    for text in pages:
        esc = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        lines = [esc[i : i + 90] for i in range(0, len(esc), 90)] or [""]
        stream = b"BT /F1 11 Tf 50 750 Td 14 TL " + b" T* ".join(b"(" + ln.encode("latin-1", "replace") + b") Tj" for ln in lines) + b" ET"
        c = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        p = add(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {c} 0 R /Resources << /Font << /F1 {font} 0 R >> >> >>".encode())
        kids.append(p)
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>".encode()
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


DEMO_USERS = [
    ("admin@demo.local", "Demo Admin", ["admin"], None),
    ("hr@demo.local", "Demo HR Manager", ["hr"], None),
    ("finance@demo.local", "Demo Finance Analyst", ["finance"], None),
    ("eng@demo.local", "Demo Engineer", ["engineering"], None),
    ("junior@demo.local", "Demo Junior Developer", ["engineering"], None),
    ("employee@demo.local", "Demo Employee", ["employee"], None),
]


def _ingest(admin: User, filename: str, data: bytes, roles: list[str], classification: str) -> Document:
    name, ftype = validate_upload(filename, data, 25 * 1024 * 1024, 200)
    doc = Document(
        original_filename=name, storage_name=pipeline.store_file(data, ftype), file_type=ftype, size_bytes=len(data),
        checksum_sha256=pipeline.sha256_hex(data), uploaded_by=admin.id, classification=classification,
    )
    db.session.add(doc)
    db.session.flush()
    for r in Role.query.filter(Role.name.in_(roles)).all():
        doc.role_grants.append(DocumentRoleGrant(document_id=doc.id, role_id=r.id, granted_by=admin.id))
    db.session.commit()
    return pipeline.ingest_document(doc.id)


def load_demo(password: str, sample_dir: Path | None = None) -> None:
    seed_roles_and_permissions()
    roles = {r.name: r for r in Role.query.all()}
    users = {}
    for email, name, role_names, dept in DEMO_USERS:
        user = User.query.filter_by(email=email).first()
        if user is None:
            user = User(email=email, full_name=name, password_hash=hash_password(password),
                        roles=[roles[r] for r in role_names], department=dept)
            db.session.add(user)
        users[email] = user
    db.session.commit()
    admin = users["admin@demo.local"]
    if Document.query.filter_by(uploaded_by=admin.id).count():
        print("Demo documents already loaded.")
        return
    docs = [
        ("hr_leave_policy.md", HR_POLICY_MD.encode(), ["hr"], "confidential"),
        ("hr_salaries.xlsx", make_xlsx(SALARY_ROWS, "Salaries"), ["hr"], "highly_confidential"),
        ("engineering_architecture.md", ENG_ARCH_MD.encode(), ["engineering"], "internal"),
        ("finance_budget.csv", make_csv(BUDGET_ROWS), ["finance"], "confidential"),
        ("company_handbook.txt", HANDBOOK_TXT.encode(), ["hr", "finance", "engineering", "employee"], "internal"),
    ]
    for filename, data, role_names, cls in docs:
        result = _ingest(admin, filename, data, role_names, cls)
        print(f"{filename}: {result.status}" + (f" ({result.error_message})" if result.error_message else ""))
    if sample_dir:  # also write the files so they can be uploaded manually through the UI
        sample_dir.mkdir(parents=True, exist_ok=True)
        for filename, data, _, _ in docs:
            (sample_dir / filename).write_bytes(data)
    print("Demo users (all share the password you supplied): " + ", ".join(u[0] for u in DEMO_USERS))
