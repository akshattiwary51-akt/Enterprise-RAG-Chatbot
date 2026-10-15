"""Format-specific extraction into `Segment`s that keep provenance (page / section / sheet+row).

Spreadsheets are treated as structured data: every row becomes its own segment of the form
`Sheet: X | Row 5 | Header1: value; Header2: value`, so column names always travel with values
and unrelated rows are never merged mid-row. Formulas are never evaluated -- only cached
values are read (`data_only=True`) -- and macros are rejected at upload time.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

from .validation import decode_text


class IngestionError(Exception):
    """Processing failure with a message that is safe to store and show."""


@dataclass
class Segment:
    kind: str  # "text" | "row"
    text: str
    locator: dict = field(default_factory=dict)


# --------------------------------------------------------------------------- helpers
_WS_RE = re.compile(r"[ \t ]+")
_BLANKS_RE = re.compile(r"\n{3,}")
_CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def normalize_text(text: str) -> str:
    """Light normalisation: keeps digits, punctuation, currency symbols and line structure."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CTRL_RE.sub("", text)
    text = _WS_RE.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _BLANKS_RE.sub("\n\n", text).strip()


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else repr(value)
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0, 0) else value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    return normalize_text(str(value)).replace("\n", " ")


def _make_headers(first_row: list[str]) -> list[str]:
    headers, seen = [], {}
    for i, h in enumerate(first_row):
        name = h.strip() or f"Column{i + 1}"
        n = seen.get(name, 0)
        seen[name] = n + 1
        headers.append(name if n == 0 else f"{name}_{n + 1}")
    return headers


def _table_segments(sheet: str, rows, row_budget: list[int]) -> list[Segment]:
    """rows: iterable of (row_number, list[str]). First non-empty row = headers."""
    segs: list[Segment] = []
    headers: list[str] | None = None
    for row_no, values in rows:
        if not any(v.strip() for v in values):
            continue
        if headers is None:
            headers = _make_headers(values)
            continue
        row_budget[0] -= 1
        if row_budget[0] < 0:
            raise IngestionError("Spreadsheet has more rows than the configured limit")
        pairs = [f"{h}: {v}" for h, v in zip(headers, values) if v.strip() != ""]
        if not pairs:
            continue
        segs.append(
            Segment("row", f"Sheet: {sheet} | Row {row_no} | " + "; ".join(pairs), {"sheet": sheet, "row": row_no})
        )
    return segs


# --------------------------------------------------------------------------- parsers
def parse_pdf(path: Path) -> list[Segment]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise IngestionError("PDF is password protected")
            except IngestionError:
                raise
            except Exception as exc:
                raise IngestionError("PDF is password protected or uses unsupported encryption") from exc
        segs = []
        for number, page in enumerate(reader.pages, start=1):
            text = normalize_text(page.extract_text() or "")
            if text:
                segs.append(Segment("text", text, {"page": number}))
    except IngestionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError, RecursionError) as exc:
        raise IngestionError("PDF could not be read (file may be corrupt)") from exc
    if not segs:
        raise IngestionError("No extractable text found in the PDF (scanned PDFs need OCR, which is not supported)")
    return segs


def parse_docx(path: Path) -> list[Segment]:
    try:
        from docx import Document as Docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Docx(str(path))
    except Exception as exc:
        raise IngestionError("Word document could not be read (file may be corrupt)") from exc

    segs: list[Segment] = []
    heading_stack: list[tuple[int, str]] = []
    buf: list[str] = []

    def flush():
        text = normalize_text("\n".join(buf))
        if text:
            title = " > ".join(h for _, h in heading_stack) or "Document"
            segs.append(Segment("text", text, {"section": title}))
        buf.clear()

    try:
        for child in doc.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                para = Paragraph(child, doc)
                text = para.text.strip()
                if not text:
                    continue
                style = (para.style.name if para.style is not None else "") or ""
                m = re.match(r"Heading (\d)", style)
                if m:
                    flush()
                    level = int(m.group(1))
                    while heading_stack and heading_stack[-1][0] >= level:
                        heading_stack.pop()
                    heading_stack.append((level, text))
                    buf.append(text)
                else:
                    buf.append(text)
            elif tag == "tbl":
                table = Table(child, doc)
                rows = [[_cell_to_str(c.text) for c in r.cells] for r in table.rows]
                if rows:
                    headers = _make_headers(rows[0])
                    for r in rows[1:]:
                        buf.append("; ".join(f"{h}: {v}" for h, v in zip(headers, r) if v))
        flush()
    except Exception as exc:
        raise IngestionError("Word document structure could not be parsed") from exc
    if not segs:
        raise IngestionError("No extractable text found in the document")
    return segs


def parse_text(path: Path, markdown: bool) -> list[Segment]:
    try:
        text = normalize_text(decode_text(path.read_bytes()))
    except Exception as exc:
        raise IngestionError("Text file could not be decoded") from exc
    if not text:
        raise IngestionError("File contains no text")
    if not markdown:
        return [Segment("text", text, {"section": "Text"})]
    segs: list[Segment] = []
    stack: list[tuple[int, str]] = []
    buf: list[str] = []
    in_fence = False

    def flush():
        body = "\n".join(buf).strip()
        if body:
            segs.append(Segment("text", body, {"section": " > ".join(h for _, h in stack) or "Document"}))
        buf.clear()

    for line in text.split("\n"):
        if line.strip().startswith("```"):
            in_fence = not in_fence
        m = None if in_fence else re.match(r"^(#{1,6})\s+(.*\S)\s*$", line)
        if m:
            flush()
            level = len(m.group(1))
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, m.group(2)))
            buf.append(line)
        else:
            buf.append(line)
    flush()
    return segs


def parse_csv(path: Path, row_limit: int) -> list[Segment]:
    try:
        raw = decode_text(path.read_bytes())
        try:
            dialect = csv.Sniffer().sniff(raw[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(raw), dialect)
        rows = ((i, [_cell_to_str(c) for c in rec]) for i, rec in enumerate(reader, start=1))
        segs = _table_segments(path.stem[:60] or "csv", rows, [row_limit])
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("CSV could not be parsed") from exc
    if not segs:
        raise IngestionError("CSV contains a header but no data rows")
    return segs


def parse_xlsx(path: Path, row_limit: int) -> list[Segment]:
    from openpyxl import load_workbook

    budget = [row_limit]
    segs: list[Segment] = []
    try:
        wb = load_workbook(str(path), read_only=True, data_only=True)  # cached values only; no formula eval
        try:
            for ws in wb.worksheets:
                if getattr(ws, "sheet_state", "visible") != "visible":
                    continue  # hidden sheets are not indexed
                rows = ((i, [_cell_to_str(c) for c in row]) for i, row in enumerate(ws.iter_rows(values_only=True), 1))
                segs.extend(_table_segments(ws.title, rows, budget))
        finally:
            wb.close()
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("Excel workbook could not be read (file may be corrupt)") from exc
    if not segs:
        raise IngestionError("Workbook contains no data rows")
    return segs


def parse_xls(path: Path, row_limit: int) -> list[Segment]:
    import pandas as pd

    budget = [row_limit]
    segs: list[Segment] = []
    try:
        sheets = pd.read_excel(str(path), sheet_name=None, header=None, dtype=object, engine="xlrd")
        for name, frame in sheets.items():
            rows = (
                (i, [_cell_to_str(None if pd.isna(v) else v) for v in rec])
                for i, rec in enumerate(frame.itertuples(index=False, name=None), start=1)
            )
            segs.extend(_table_segments(str(name), rows, budget))
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("Legacy Excel file could not be read (file may be corrupt)") from exc
    if not segs:
        raise IngestionError("Workbook contains no data rows")
    return segs


def iter_sheet_rows(path: Path, file_type: str):
    """Yield (sheet_name, [(row_no, [str, ...]), ...]) for tabular files. Used by the
    structured-query engine so it sees exactly the values that were indexed."""
    if file_type == "csv":
        raw = decode_text(path.read_bytes())
        try:
            dialect = csv.Sniffer().sniff(raw[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        yield path.stem[:60] or "csv", [
            (i, [_cell_to_str(c) for c in rec]) for i, rec in enumerate(csv.reader(io.StringIO(raw), dialect), 1)
        ]
    elif file_type == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                if getattr(ws, "sheet_state", "visible") != "visible":
                    continue
                yield ws.title, [
                    (i, [_cell_to_str(c) for c in row]) for i, row in enumerate(ws.iter_rows(values_only=True), 1)
                ]
        finally:
            wb.close()
    elif file_type == "xls":
        import pandas as pd

        for name, frame in pd.read_excel(str(path), sheet_name=None, header=None, dtype=object, engine="xlrd").items():
            yield str(name), [
                (i, [_cell_to_str(None if pd.isna(v) else v) for v in rec])
                for i, rec in enumerate(frame.itertuples(index=False, name=None), start=1)
            ]


def parse_file(path: Path, file_type: str, row_limit: int) -> list[Segment]:
    if file_type == "pdf":
        return parse_pdf(path)
    if file_type == "docx":
        return parse_docx(path)
    if file_type in ("txt", "md"):
        return parse_text(path, markdown=file_type == "md")
    if file_type == "csv":
        return parse_csv(path, row_limit)
    if file_type == "xlsx":
        return parse_xlsx(path, row_limit)
    if file_type == "xls":
        return parse_xls(path, row_limit)
    raise IngestionError("Unsupported file type")
